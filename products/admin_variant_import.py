
import os
from decimal import Decimal, InvalidOperation
from urllib.parse import urlparse

import requests
from django.contrib import admin, messages
from django.core.files.base import ContentFile
from django.shortcuts import render, redirect
from django.urls import path

from .models import Product, ProductVariant


REQUIRED_COLUMNS = {
    "product_name", "color", "color_code", "storage",
    "image_url", "price", "discount_price", "stock", "sku", "is_active"
}


def _clean(value):
    if value is None:
        return ""
    return str(value).strip()


def _decimal(value):
    value = _clean(value)
    if not value:
        return None
    try:
        return Decimal(value)
    except (InvalidOperation, ValueError):
        raise ValueError(f"Invalid price: {value}")


def _int(value, default=0):
    value = _clean(value)
    if not value:
        return default
    try:
        return int(float(value))
    except (ValueError, TypeError):
        raise ValueError(f"Invalid stock: {value}")


def _bool(value):
    value = _clean(value).lower()
    return value not in {"0", "false", "no", "off", "inactive"}


def _download_image(url, variant):
    url = _clean(url)
    if not url:
        return False

    parsed = urlparse(url)
    if parsed.scheme not in {"http", "https"}:
        raise ValueError("image_url must start with http:// or https://")

    response = requests.get(
        url,
        timeout=20,
        headers={"User-Agent": "ShopEase Excel Importer/1.0"},
    )
    response.raise_for_status()

    content_type = response.headers.get("Content-Type", "").lower()
    if not content_type.startswith("image/"):
        raise ValueError("image_url did not return an image")

    filename = os.path.basename(parsed.path) or f"variant-{variant.pk}.jpg"
    variant.image.save(filename, ContentFile(response.content), save=True)
    return True


class ExcelVariantImportMixin:
    """
    Add this mixin before the existing ProductVariantAdmin class.

    Example:
        @admin.register(ProductVariant)
        class ProductVariantAdmin(ExcelVariantImportMixin, admin.ModelAdmin):
            ...
    """

    def get_urls(self):
        urls = super().get_urls()
        custom = [
            path(
                "import-excel/",
                self.admin_site.admin_view(self.import_excel_view),
                name="products_productvariant_import_excel",
            ),
        ]
        return custom + urls

    def import_excel_view(self, request):
        if request.method == "POST":
            uploaded = request.FILES.get("excel_file")

            if not uploaded:
                messages.error(request, "Excel fayl tanlang.")
                return redirect("admin:products_productvariant_import_excel")

            if not uploaded.name.lower().endswith((".xlsx", ".xlsm")):
                messages.error(request, "Faqat .xlsx yoki .xlsm fayl yuklang.")
                return redirect("admin:products_productvariant_import_excel")

            try:
                from openpyxl import load_workbook
                wb = load_workbook(uploaded, data_only=True)
                ws = wb.active

                headers = [_clean(c.value) for c in ws[1]]
                missing = REQUIRED_COLUMNS - set(headers)
                if missing:
                    messages.error(
                        request,
                        "Excel ustunlari yetishmayapti: " + ", ".join(sorted(missing))
                    )
                    return redirect("admin:products_productvariant_import_excel")

                index = {name: i for i, name in enumerate(headers)}

                created = 0
                updated = 0
                images = 0
                errors = []

                for row_number, row in enumerate(
                    ws.iter_rows(min_row=2, values_only=True),
                    start=2,
                ):
                    if not any(v not in (None, "") for v in row):
                        continue

                    try:
                        product_name = _clean(row[index["product_name"]])
                        color = _clean(row[index["color"]])
                        color_code = _clean(row[index["color_code"]])
                        storage = _clean(row[index["storage"]])
                        image_url = _clean(row[index["image_url"]])
                        sku = _clean(row[index["sku"]])

                        if not product_name:
                            raise ValueError("product_name bo'sh")
                        if not color:
                            raise ValueError("color bo'sh")
                        if not storage:
                            raise ValueError("storage bo'sh")
                        if not sku:
                            raise ValueError("sku bo'sh")

                        product = Product.objects.filter(
                            name__iexact=product_name
                        ).first()

                        if not product:
                            raise ValueError(
                                f"Product topilmadi: {product_name}"
                            )

                        defaults = {
                            "color_code": color_code,
                            "storage": storage,
                            "price": _decimal(row[index["price"]]),
                            "discount_price": _decimal(
                                row[index["discount_price"]]
                            ),
                            "stock": _int(row[index["stock"]]),
                            "is_active": _bool(row[index["is_active"]]),
                        }

                        variant, was_created = ProductVariant.objects.update_or_create(
                            product=product,
                            color=color,
                            storage=storage,
                            defaults={**defaults, "sku": sku},
                        )

                        if was_created:
                            created += 1
                        else:
                            updated += 1

                        if image_url:
                            if _download_image(image_url, variant):
                                images += 1

                    except Exception as exc:
                        errors.append(f"{row_number}-qator: {exc}")

                if errors:
                    messages.warning(
                        request,
                        f"Import tugadi. Yaratildi: {created}, yangilandi: {updated}, "
                        f"rasm yuklandi: {images}. Xatolar: {len(errors)}. "
                        + " | ".join(errors[:5])
                    )
                else:
                    messages.success(
                        request,
                        f"Import muvaffaqiyatli. Yaratildi: {created}, "
                        f"yangilandi: {updated}, rasmlar: {images}."
                    )

                return redirect("admin:products_productvariant_changelist")

            except Exception as exc:
                messages.error(request, f"Import xatosi: {exc}")
                return redirect("admin:products_productvariant_import_excel")

        context = {
            **self.admin_site.each_context(request),
            "title": "Import product variants from Excel",
            "opts": self.model._meta,
        }
        return render(
            request,
            "admin/products/productvariant/import_excel.html",
            context,
        )
